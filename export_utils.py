"""
邮件查询结果导出工具模块。

提供将查询结果导出为 Excel (.xlsx) 和 CSV 格式的功能。
"""

import csv
import os
import tempfile
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def export_to_excel(
    data: list[dict[str, Any]],
    filepath: str,
    columns: Optional[list[str]] = None,
    sheet_title: str = "邮件查询结果",
) -> str:
    """将数据导出为 Excel 文件。

    Args:
        data: 数据行列表，每行为一个 dict，key 为列名，value 为单元格值。
        filepath: 输出文件的绝对路径。
        columns: 列名列表，决定导出列的顺序。为 None 时取 data[0].keys()。
        sheet_title: 工作表标题。

    Returns:
        成功时返回输出文件路径。

    Raises:
        ValueError: data 为 None 或空列表。
        PermissionError: 文件被占用无法写入。
        OSError: 磁盘满或其他 IO 错误。
    """
    if data is None or len(data) == 0:
        raise ValueError("data 不能为 None 或空列表")

    if columns is None:
        columns = list(data[0].keys())

    raw_path = Path(filepath)
    filename = raw_path.name
    dest_dir = raw_path.parent.resolve()
    if not dest_dir.exists():
        dest_dir.mkdir(parents=True, exist_ok=True)
    dest_path = dest_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    # 写入表头
    for col_idx, col_name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    # 写入数据行
    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ""))

    # 应用格式
    num_rows = len(data) + 1
    num_cols = len(columns)
    _apply_excel_format(ws, num_rows, num_cols)
    _auto_column_width(ws)

    try:
        fd, tmp_path = tempfile.mkstemp(dir=str(dest_dir), suffix=".xlsx")
        os.close(fd)
        wb.save(tmp_path)
        os.replace(tmp_path, str(dest_path))
    except PermissionError:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise PermissionError(f"无法写入文件，文件可能已被其他程序占用: {dest_path.name}")
    except OSError as e:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise OSError(f"无法写入文件，请检查磁盘空间: {e}")

    return str(dest_path)


def _apply_excel_format(worksheet: Worksheet, num_rows: int, num_cols: int) -> None:
    """对工作表应用统一的格式化样式。

    Args:
        worksheet: 目标工作表对象。
        num_rows: 数据总行数（含表头）。
        num_cols: 数据总列数。
    """
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_alignment = Alignment(vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # 表头行格式化
    for col_idx in range(1, num_cols + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 设置表头行高
    worksheet.row_dimensions[1].height = 28

    # 数据行格式化
    for row_idx in range(2, num_rows + 1):
        for col_idx in range(1, num_cols + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = data_alignment
            cell.border = thin_border

    # 冻结首行
    worksheet.freeze_panes = "A2"

    # 添加自动筛选
    worksheet.auto_filter.ref = f"A1:{get_column_letter(num_cols)}{num_rows}"


def _auto_column_width(
    worksheet: Worksheet,
    min_width: int = 8,
    max_width: int = 50,
) -> None:
    """根据单元格内容自动计算并设置列宽。

    中文字符计为 2 单位宽度，ASCII 字符计为 1 单位宽度。

    Args:
        worksheet: 目标工作表对象。
        min_width: 最小列宽。
        max_width: 最大列宽。
    """
    for col_cells in worksheet.columns:
        max_char_width = 0
        col_letter = get_column_letter(col_cells[0].column)

        for cell in col_cells:
            if cell.value is None:
                continue
            cell_str = str(cell.value)
            char_width = 0
            for ch in cell_str:
                if ord(ch) > 127:
                    char_width += 2
                else:
                    char_width += 1
            if char_width > max_char_width:
                max_char_width = char_width

        # clamp 到指定范围，并加一点 padding
        adjusted_width = max(min_width, min(max_char_width + 2, max_width))
        worksheet.column_dimensions[col_letter].width = adjusted_width


def export_csv(
    data: list[dict[str, Any]],
    filepath: str,
    columns: Optional[list[str]] = None,
) -> str:
    """将数据导出为 CSV 文件。

    使用 utf-8-sig 编码，确保 Excel 能正确识别中文。

    Args:
        data: 数据行列表，每行为一个 dict，key 为列名，value 为单元格值。
        filepath: 输出文件的绝对路径。
        columns: 列名列表，决定导出列的顺序。为 None 时取 data[0].keys()。

    Returns:
        成功时返回输出文件路径。

    Raises:
        ValueError: data 为 None 或空列表。
        OSError: 磁盘满或其他 IO 错误。
    """
    if data is None or len(data) == 0:
        raise ValueError("data 不能为 None 或空列表")

    if columns is None:
        columns = list(data[0].keys())

    raw_path = Path(filepath)
    filename = raw_path.name
    dest_dir = raw_path.parent.resolve()
    dest_path = dest_dir / filename

    try:
        fd, tmp_path = tempfile.mkstemp(dir=str(dest_path.parent), suffix=".csv")
        with os.fdopen(fd, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            for row_data in data:
                writer.writerow([row_data.get(col, "") for col in columns])
        os.replace(tmp_path, str(dest_path))
    except PermissionError:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise PermissionError(f"无法写入文件，文件可能已被其他程序占用: {dest_path.name}")
    except OSError as e:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        raise OSError(f"无法写入文件，请检查磁盘空间: {e}")

    return str(dest_path)